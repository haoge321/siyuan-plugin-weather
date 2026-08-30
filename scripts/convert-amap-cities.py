import argparse
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_HEADERS = ("中文名", "adcode", "citycode")
ROOT_ADCODE = "100000"


def normalize_cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_source(source_path):
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError(
                f"expected exactly one worksheet, found {len(workbook.worksheets)}"
            )

        worksheet = workbook.active
        header = tuple(
            normalize_cell(value)
            for value in next(worksheet.iter_rows(values_only=True))
        )[:3]
        if header != EXPECTED_HEADERS:
            raise ValueError(
                f"unexpected header: {header!r}; expected {EXPECTED_HEADERS!r}"
            )

        records = []
        for source_row, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            values = tuple(normalize_cell(value) for value in row[:3])
            if not any(values):
                continue
            if len(values) != 3 or any(not value for value in values):
                raise ValueError(
                    f"incomplete data at Excel row {source_row}: {values!r}"
                )

            name, adcode, citycode = values
            if len(adcode) != 6 or not adcode.isdigit():
                raise ValueError(
                    f"invalid adcode at Excel row {source_row}: {adcode!r}"
                )

            records.append(
                {
                    "name": name,
                    "adcode": adcode,
                    "citycode": citycode,
                    "sourceRow": source_row,
                }
            )

        if not records:
            raise ValueError("source workbook contains no data rows")
        return records
    finally:
        workbook.close()


def get_level(adcode):
    if adcode == ROOT_ADCODE:
        return "country"
    if adcode.endswith("0000"):
        return "province"
    if adcode.endswith("00"):
        return "city"
    return "district"


def build_records(source_records):
    by_adcode = {record["adcode"]: record for record in source_records}
    if len(by_adcode) != len(source_records):
        duplicates = [
            adcode
            for adcode, count in Counter(
                record["adcode"] for record in source_records
            ).items()
            if count > 1
        ]
        raise ValueError(f"duplicate adcode values: {duplicates[:10]}")

    root = by_adcode.get(ROOT_ADCODE)
    if root is None:
        raise ValueError(f"missing root record {ROOT_ADCODE}")

    path_by_adcode = {}
    output = []
    for source_record in source_records:
        adcode = source_record["adcode"]
        level = get_level(adcode)

        if level == "country":
            parent_adcode = None
        elif level == "province":
            parent_adcode = ROOT_ADCODE
        elif level == "city":
            parent_adcode = f"{adcode[:2]}0000"
        else:
            city_parent = f"{adcode[:4]}00"
            province_parent = f"{adcode[:2]}0000"
            parent_adcode = (
                city_parent
                if city_parent in by_adcode
                else province_parent
            )

        if parent_adcode is not None and parent_adcode not in by_adcode:
            raise ValueError(
                f"missing parent {parent_adcode} for {adcode} ({source_record['name']})"
            )

        if parent_adcode is None:
            path = source_record["name"]
        else:
            parent_path = path_by_adcode.get(parent_adcode)
            if parent_path is None:
                raise ValueError(
                    f"parent {parent_adcode} must appear before child {adcode}"
                )
            path = f"{parent_path} / {source_record['name']}"

        converted = {
            "name": source_record["name"],
            "adcode": adcode,
            "citycode": source_record["citycode"],
            "level": level,
            "parentAdcode": parent_adcode,
            "path": path,
            "selectable": level != "country",
        }
        path_by_adcode[adcode] = path
        output.append(converted)

    if output[0]["adcode"] != ROOT_ADCODE:
        raise ValueError("the first data record must be the country root")

    return output


def compact_records(output):
    if not output or output[0]["adcode"] != ROOT_ADCODE:
        raise ValueError("cannot compact output without the country root")

    compacted = []
    for record in output[1:]:
        parent_adcode = record["parentAdcode"]
        compacted.append(
            [
                record["name"],
                record["adcode"],
                None if parent_adcode == ROOT_ADCODE else parent_adcode,
            ]
        )
    return compacted


def validate_compact_output(output, compacted):
    if len(compacted) != len(output) - 1:
        raise ValueError(
            f"compact record count changed: output={len(output)}, "
            f"compact={len(compacted)}"
        )

    output_identity = [
        (record["name"], record["adcode"], record["parentAdcode"])
        for record in output[1:]
    ]
    compact_identity = [
        (record[0], record[1], ROOT_ADCODE if record[2] is None else record[2])
        for record in compacted
    ]
    if output_identity != compact_identity:
        raise ValueError("compact name/adcode/parent values changed")

    compact_adcodes = [record[1] for record in compacted]
    if len(compact_adcodes) != len(set(compact_adcodes)):
        raise ValueError("duplicate adcode values in compact output")

    available_adcodes = set(compact_adcodes)
    for record in compacted:
        parent_adcode = record[2]
        if parent_adcode is not None and parent_adcode not in available_adcodes:
            raise ValueError(
                f"compact output parent {parent_adcode} not found for {record[1]}"
            )


def validate_output(source_records, output):
    if len(source_records) != len(output):
        raise ValueError(
            f"record count changed: source={len(source_records)}, output={len(output)}"
        )

    source_identity = [
        (record["name"], record["adcode"], record["citycode"])
        for record in source_records
    ]
    output_identity = [
        (record["name"], record["adcode"], record["citycode"])
        for record in output
    ]
    if source_identity != output_identity:
        raise ValueError("name/adcode/citycode values or source order changed")

    by_adcode = {record["adcode"]: record for record in output}
    for record in output:
        parent_adcode = record["parentAdcode"]
        if parent_adcode is None:
            if record["level"] != "country" or record["path"] != record["name"]:
                raise ValueError(f"invalid root record: {record!r}")
            continue

        parent = by_adcode.get(parent_adcode)
        if parent is None:
            raise ValueError(
                f"output parent {parent_adcode} not found for {record['adcode']}"
            )
        expected_path = f"{parent['path']} / {record['name']}"
        if record["path"] != expected_path:
            raise ValueError(
                f"invalid path for {record['adcode']}: {record['path']!r}"
            )

        visited = {record["adcode"]}
        current = parent
        while current["parentAdcode"] is not None:
            if current["adcode"] in visited:
                raise ValueError(f"parent cycle detected at {current['adcode']}")
            visited.add(current["adcode"])
            current = by_adcode[current["parentAdcode"]]

    if sum(record["selectable"] for record in output) != len(output) - 1:
        raise ValueError("only the country root should be non-selectable")


def parse_args():
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(
        description="Convert the Amap weather city code workbook to validated JSON."
    )
    parser.add_argument(
        "--source",
        type=Path,
        default=project_root / "doc" / "高德天气城市编码对照表.xlsx",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=project_root / "src" / "data" / "amap-cities.json",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    source_path = args.source.resolve()
    output_path = args.output.resolve()

    if not source_path.is_file():
        raise FileNotFoundError(f"source workbook not found: {source_path}")

    source_records = read_source(source_path)
    output = build_records(source_records)
    validate_output(source_records, output)
    compacted_output = compact_records(output)
    validate_compact_output(output, compacted_output)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(compacted_output, ensure_ascii=False, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )

    level_counts = Counter(record["level"] for record in output)
    duplicate_names = len(output) - len({record["name"] for record in output})
    print(f"Converted {len(output)} records.")
    print(f"Compact records: {len(compacted_output)} (country root omitted).")
    print(f"Levels: {dict(level_counts)}")
    print(f"Duplicate names kept for disambiguation: {duplicate_names}")
    print(f"Output: {output_path}")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"Conversion failed: {error}", file=sys.stderr)
        sys.exit(1)
